"""Append dated progress bullets to the existing RecruitO project log."""

import argparse
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document


LOG_PATH = Path(__file__).with_name("RecruitO_Project_Progress_Log.docx")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--date",
        default=datetime.now(ZoneInfo("Asia/Dhaka")).strftime("%d %B %Y %H:%M"),
    )
    parser.add_argument("--item", action="append", help="Progress point; repeat for multiple points")
    parser.add_argument("--replace", action="append", default=[], metavar="OLD|||NEW", help="Replace stale wording in an existing paragraph")
    args = parser.parse_args()
    if not args.item and not args.replace:
        parser.error("provide at least one --item or --replace")

    document = Document(LOG_PATH)
    for replacement in args.replace:
        old, separator, new = replacement.partition("|||")
        if not separator:
            parser.error("--replace must use OLD|||NEW")
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                updated_text = paragraph.text.replace(old, new)
                for run in paragraph.runs:
                    run.text = ""
                paragraph.add_run(updated_text)
    updates_heading = next(
        (p for p in document.paragraphs if p.style.name == "Heading 1" and p.text == "Incremental Updates"),
        None,
    )
    if updates_heading is None:
        updates_heading = document.add_heading("Incremental Updates", level=1)

    document.add_heading(args.date, level=2)
    for item in args.item or []:
        bullet = document.add_paragraph(style="List Bullet")
        bullet.paragraph_format.space_after = 0
        bullet.add_run(item)

    document.save(LOG_PATH)
    print(LOG_PATH)


if __name__ == "__main__":
    main()
